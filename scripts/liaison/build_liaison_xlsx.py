#!/usr/bin/env python3
"""연계 정리 MD(탭 구분 코드블록) → 세명소프트 양식과 같은 모양의 엑셀.

    python scripts/liaison/build_liaison_xlsx.py            # docs/liaison/연계정리_이노뎁_2세부.xlsx

탭 6개: 기본정보 · 컨테이너·포트 · API 목록 · 프록시 경로 · 결과물 연계(요소기술 연계 10열) · 이슈&요청사항.
머리글은 남색 바탕 흰 글씨, 작성 기관 열은 연노랑, 연계 탭은 연계 1건 단위로 병합한다.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs/liaison/연계정리_이노뎁_2세부_작성본.md"
OUT = ROOT / "docs/liaison/연계정리_이노뎁_2세부.xlsx"

NAVY = PatternFill("solid", fgColor="1F3864")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
TITLE = PatternFill("solid", fgColor="D9E2F3")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def blocks(md: str) -> dict[str, list[list[str]]]:
    """'## N. 제목' 또는 '### N-N. 제목' 아래 첫 번째 ```이노뎁\t…``` 블록을 제목 키로 모은다."""

    out: dict[str, list[list[str]]] = {}
    for m in re.finditer(r"^(#{2,3}) ([0-9]+(?:-[A-Z0-9]+)?)\.[^\n]*\n(.*?)(?=^#{2,3} |\Z)", md, re.S | re.M):
        key = m.group(2)
        code = re.search(r"```\n(이노뎁\t.*?)```", m.group(3), re.S)
        if code:
            out[key] = [row.split("\t") for row in code.group(1).rstrip("\n").split("\n")]
    return out


def write_table(ws, start_row: int, title: str, headers: list[str], rows: list[list[str]],
                widths: list[int], merge_groups: bool = False, yellow_col: int = 1) -> int:
    """표 하나를 쓰고 다음 빈 행 번호를 돌려준다."""

    ncol = len(headers)
    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, color="1F3864", size=11)
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=ncol)
    for c in range(1, ncol + 1):
        ws.cell(row=start_row, column=c).fill = TITLE
    hdr = start_row + 1
    for c, name in enumerate(headers, 1):
        cell = ws.cell(row=hdr, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = NAVY; cell.alignment = CENTER; cell.border = BORDER
    r = hdr + 1
    group_start = None
    for row in rows:
        row = (row + [""] * ncol)[:ncol]
        first = bool(row[0].strip())
        if merge_groups and first and group_start is not None and r - 1 > group_start:
            _merge_group(ws, group_start, r - 1, ncol)
        if first:
            group_start = r
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value if value != "" else None)
            cell.alignment = WRAP; cell.border = BORDER
            if c == yellow_col and (first or not merge_groups):
                cell.fill = YELLOW
        r += 1
    if merge_groups and group_start is not None and r - 1 > group_start:
        _merge_group(ws, group_start, r - 1, ncol)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return r + 1


def _merge_group(ws, top: int, bottom: int, ncol: int) -> None:
    """연계 1건: 1~5열과 9~10열을 세로 병합 (6~8열 구분·항목·형식은 행마다)."""

    for c in list(range(1, 6)) + list(range(9, ncol + 1)):
        ws.merge_cells(start_row=top, start_column=c, end_row=bottom, end_column=c)
        ws.cell(row=top, column=c).alignment = Alignment(wrap_text=True, vertical="center")
        if c == 1:
            for rr in range(top, bottom + 1):
                ws.cell(row=rr, column=c).fill = YELLOW


def main() -> int:
    md = SRC.read_text(encoding="utf-8")
    b = blocks(md)
    need = ["1-1", "1-2", "2", "3", "4", "5-B", "6"]
    missing = [k for k in need if k not in b]
    if missing:
        print(f"MD 에서 블록을 못 찾음: {missing}", file=sys.stderr); return 1

    wb = Workbook()
    ws = wb.active; ws.title = "기본정보"
    nxt = write_table(ws, 1, "기관 구성·자산 — 기관당 1행", ["작성 기관", "담당 역할", "컨테이너", "비고"], b["1-1"], [14, 44, 40, 48])
    write_table(ws, nxt, "구동 서버 기본 정보 — 항목당 1행", ["작성 기관", "항목", "값", "비고"], b["1-2"], [14, 44, 40, 48])

    ws = wb.create_sheet("컨테이너·포트")
    write_table(ws, 1, "컨테이너·포트 — 컨테이너당 1행", ["작성 기관", "컨테이너", "이미지", "호스트 포트 → 컨테이너 포트", "역할", "비고"], b["2"], [14, 22, 20, 24, 36, 60])

    ws = wb.create_sheet("API 목록")
    write_table(ws, 1, "API 목록 — 엔드포인트당 1행", ["작성 기관", "서비스 · 포트", "메서드", "경로", "설명", "프록시 Path 경로", "비고"], b["3"], [14, 18, 9, 58, 40, 58, 44])

    ws = wb.create_sheet("프록시 경로")
    write_table(ws, 1, "프록시 경로 — 프록시 location당 1행", ["작성 기관", "요청 경로 prefix", "업스트림", "prefix 처리", "용도", "예시", "비고"], b["4"], [14, 30, 26, 11, 20, 46, 40])

    ws = wb.create_sheet("결과물 연계")
    write_table(ws, 1, "요소기술 연계 — 수신하는 것 · 제공하는 것. 주고받는 데이터와 API는 항목마다 1행",
                ["작성 기관", "방향", "연계 내용", "요소기술", "docker 통합 후 목표안", "구분", "항목", "형식 · 설명", "필요 사항", "비고"],
                b["5-B"], [12, 18, 30, 30, 36, 7, 28, 44, 40, 28], merge_groups=True)

    ws = wb.create_sheet("이슈&요청사항")
    write_table(ws, 1, "이슈&요청사항 — 다른 기관에 요청하거나 함께 정할 일당 1행",
                ["작성 기관", "구분", "항목", "내용", "영향", "필요 조치", "요청 대상", "상태"], b["6"], [12, 8, 26, 56, 32, 40, 12, 10])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A3"
        sheet.sheet_view.zoomScale = 90
    wb.save(OUT)
    counts = {k: len(v) for k, v in b.items() if k in need}
    print(f"저장: {OUT}  행수 {counts}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
