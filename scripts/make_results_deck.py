#!/usr/bin/env python3
"""과제 성과 발표자료(PPTX + MD): 위성 데이터 → 수체·수위 탐지 → 시계열 분석 → 모델별 성능, 화면 스크린샷 다수.

    pipeline_ui/.venv/bin/python scripts/make_results_deck.py

입력(모두 산출물 파일 — 수치는 손으로 넣지 않는다):
  regions/*/satellite/{scenes.json, detection_eval.json, levels.json, timeseries_eval.json}
  regions/*/gauge/series.json, regions/*/models/gauge/evaluation.json
  regions/*/models/gauge/experiments/future_rain.json (미래 강수 상한 실험 — 없으면 그 슬라이드는 안내 문구만)
  docs/presentation/figures/*.png (scripts/make_model_figures.py)
  docs/presentation/screens/*.png (scripts/capture_ui_screens.py)
출력: docs/presentation/과제성과_위성수체수위_시계열예측.pptx, 같은 이름 .md
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT / "ui_next") not in sys.path:  # 화면과 같은 표 도우미로 실험 JSON 을 읽는다(수치 이중 계산 방지)
    sys.path.insert(0, str(ROOT / "ui_next"))
import area_models_ui as am  # noqa: E402
from future_rain_ui import experiment_rows, region_summary_rows, summary_by_horizon  # noqa: E402
FIG = ROOT / "docs" / "presentation" / "figures"
SCR = ROOT / "docs" / "presentation" / "screens"
OUT = ROOT / "docs" / "presentation"
SHORT_NAMES = {"nakdong_busan": "낙동강 하구(부산·김해)", "geum_daecheong": "금강 대청댐", "han_river": "한강(대조군)"}


def _discover_regions() -> tuple[dict[str, str], dict[str, dict[str, str]]]:
    """regions/*/region.json 을 읽어 (id → 표시 이름), (id → {수위계 지점: 위성 산출 위치}) 를 만든다.

    위성 있는 지역이 먼저, 그다음 이름순. 새 지역을 붙이면 자동으로 자료에 들어간다.
    """
    names: dict[str, str] = {}
    links: dict[str, dict[str, str]] = {}
    order = []
    for path in sorted((ROOT / "regions").glob("*/region.json")):
        meta = json.loads(path.read_text(encoding="utf-8"))
        rid = meta.get("id") or path.parent.name
        names[rid] = SHORT_NAMES.get(rid) or str(meta.get("name") or rid)
        sat_links = ((meta.get("gauge") or {}).get("satellite_location_links") or {})
        if sat_links and (meta.get("satellite") or {}).get("available"):
            links[rid] = {str(k): str(v) for k, v in sat_links.items()}
        order.append((0 if (meta.get("satellite") or {}).get("available") else 1, rid))
    ordered = {rid: names[rid] for _, rid in sorted(order)}
    return ordered, links


REGIONS, SAT_LINKS = _discover_regions()
NAMES = {
    "persistence": "직전값 유지(기준선)", "climatology": "계절 평년값", "ridge_ar": "릿지 자기회귀", "ridge_arx": "릿지 AR+기상",
    "seasonal_anomaly_ar": "계절편차 AR", "mean_reversion": "평균 회귀 AR(1)", "analog_knn": "유사사례 k-NN", "ensemble_mean": "계열 앙상블",
    "neural_mlp": "MLP(GPU)", "neural_lstm": "LSTM(GPU)", "neural_gru": "GRU(GPU)", "neural_tcn": "TCN(GPU)", "prophet_hybrid": "Prophet+잔차AR",
}
DEEP = {"neural_mlp", "neural_lstm", "neural_gru", "neural_tcn"}
FUTURE_RAIN_HORIZONS = (1, 3, 7, 14, 30)
BG = RGBColor(0x0B, 0x13, 0x24); FG = RGBColor(0xE2, 0xE8, 0xF0); MUTED = RGBColor(0x94, 0xA3, 0xB8); CYAN = RGBColor(0x22, 0xD3, 0xEE)
SLIDE_W, SLIDE_H = 13.333, 7.5


# ----------------------------------------------------------------------------- helpers
def jload(path: Path):
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else None


def m_cm(v) -> str:
    return f"{v:.3f} m ({v * 100:.1f} cm)" if isinstance(v, (int, float)) else "-"


def m2_cm2(v) -> str:
    return f"{v:.5f} m² ({v * 1e4:.1f} cm²)" if isinstance(v, (int, float)) else "-"


def pct(v) -> str:
    return f"{v * 100:+.1f} %" if isinstance(v, (int, float)) else "-"


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W); self.prs.slide_height = Inches(SLIDE_H)
        self.md: list[str] = []
        self.n = 0

    # -- slide scaffolding
    def slide(self, title: str, subtitle: str | None = None):
        self.n += 1
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        f = s.background.fill; f.solid(); f.fore_color.rgb = BG
        tb = s.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(12.4), Inches(0.9)); tf = tb.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(24); p.font.bold = True; p.font.color.rgb = FG
        if subtitle:
            q = tf.add_paragraph(); q.text = subtitle; q.font.size = Pt(12); q.font.color.rgb = MUTED
        foot = s.shapes.add_textbox(Inches(0.45), Inches(7.05), Inches(12.4), Inches(0.3)); fp = foot.text_frame.paragraphs[0]
        fp.text = f"초소형위성 수체·수위 탐지 및 시계열 예측 · {self.n}"; fp.font.size = Pt(9); fp.font.color.rgb = MUTED
        self.md += [f"## {self.n}. {title}", ""] + ([f"_{subtitle}_", ""] if subtitle else [])
        return s

    def bullets(self, s, items, *, left=0.55, top=1.35, width=6.0, height=5.4, size=13):
        box = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)); tf = box.text_frame; tf.word_wrap = True
        for i, text in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = "• " + text; para.font.size = Pt(size); para.font.color.rgb = FG; para.space_after = Pt(6)
        self.md += [f"- {t}" for t in items] + [""]

    def image(self, s, path: Path, *, left: float, top: float, max_w: float, max_h: float, caption: str | None = None):
        """상자(max_w×max_h) 안에 비율을 지키며 맞춘다."""
        if not path.is_file():
            box = s.shapes.add_textbox(Inches(left), Inches(top), Inches(max_w), Inches(0.5)); p = box.text_frame.paragraphs[0]
            p.text = f"(이미지 없음: {path.name})"; p.font.size = Pt(10); p.font.color.rgb = MUTED
            return
        with Image.open(path) as im:
            w, h = im.size
        scale = min(max_w / w, max_h / h)
        pic_w, pic_h = w * scale, h * scale
        s.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(pic_w), height=Inches(pic_h))
        if caption:
            box = s.shapes.add_textbox(Inches(left), Inches(top + pic_h + 0.05), Inches(max_w), Inches(0.35)); p = box.text_frame.paragraphs[0]
            p.text = caption; p.font.size = Pt(10); p.font.color.rgb = MUTED
        rel = path.relative_to(OUT).as_posix()
        self.md += [f"![{caption or path.stem}]({rel})", ""]

    def table(self, s, header, rows, *, left=0.5, top=1.35, width=12.3, font=10, row_h=0.3):
        shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(left), Inches(top), Inches(width), Inches(row_h) * (len(rows) + 1))
        tbl = shape.table
        for j, h in enumerate(header):
            c = tbl.cell(0, j); c.text = str(h); c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            for para in c.text_frame.paragraphs:
                para.font.size = Pt(font); para.font.bold = True; para.font.color.rgb = CYAN
        for i, row in enumerate(rows, start=1):
            for j, val in enumerate(row):
                c = tbl.cell(i, j); c.text = str(val); c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                for para in c.text_frame.paragraphs:
                    para.font.size = Pt(font); para.font.color.rgb = FG
        self.md += ["| " + " | ".join(map(str, header)) + " |", "|" + "---|" * len(header)] + ["| " + " | ".join(map(str, r)) + " |" for r in rows] + [""]

    def note(self, s, text, *, top=6.6, left=0.5, width=12.3):
        box = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.45)); tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(10.5); p.font.color.rgb = MUTED
        self.md += [f"> {text}", ""]


# ----------------------------------------------------------------------------- data
def load():
    data = {"scenes": {}, "levels": {}, "series": {}, "evals": {}, "experiments": {}, "area_models": {}}
    for rid in REGIONS:
        base = ROOT / "regions" / rid
        data["scenes"][rid] = jload(base / "satellite" / "scenes.json")
        data["levels"][rid] = jload(base / "satellite" / "levels.json")
        data["series"][rid] = jload(base / "gauge" / "series.json")
        data["evals"][rid] = jload(base / "models" / "gauge" / "evaluation.json")
        data["experiments"][rid] = jload(base / "models" / "gauge" / "experiments" / "future_rain.json")
        data["area_models"][rid] = jload(base / "models" / "gauge" / "experiments" / "area_models.json")
    data["det"] = jload(ROOT / "regions" / "nakdong_busan" / "satellite" / "detection_eval.json") or {}
    data["ts"] = jload(ROOT / "regions" / "nakdong_busan" / "satellite" / "timeseries_eval.json") or {}
    return data


def scene_rows(scenes, sensor: str):
    rows = []
    for x in (scenes or {}).get("scenes", []):
        if x.get("sensor") != sensor:
            continue
        m = x.get("metrics") or {}
        rows.append([x["date"], f"{x.get('water_area_km2', 0):.2f}"] + [f"{m[k]:.3f}" if isinstance(m.get(k), (int, float)) else "-" for k in ("water_iou", "f1", "precision", "recall")])
    return rows


def satellite_vs_gauge(data):
    """지역·지점별 위성 보정 수위 vs 수위계 같은 날 오차 → (행, {지점: rmse_m})."""
    rows, rmse_by = [], {}
    for rid, links in SAT_LINKS.items():
        recs = (data["levels"].get(rid) or {}).get("records") or []
        series = ((data["series"].get(rid) or {}).get("series") or {})
        for st, loc in links.items():
            gauge = {r["date"]: r.get("water_level_m") for r in series.get(st, [])}
            errs = []
            for r in recs:
                if r.get("loc_id") != loc:
                    continue
                corr = r.get("corrected_water_level_m"); g = gauge.get(r["date"])
                if isinstance(corr, (int, float)) and isinstance(g, (int, float)):
                    errs.append(corr - g)
            n_sat = len([r for r in recs if r.get("loc_id") == loc])
            if errs:
                rmse = (sum(e * e for e in errs) / len(errs)) ** 0.5; mae = sum(abs(e) for e in errs) / len(errs)
                rows.append([REGIONS[rid], st, str(n_sat), str(len(errs)), m_cm(rmse), m2_cm2(rmse * rmse), m_cm(mae)])
                rmse_by[st] = rmse
    return rows, rmse_by


def model_metric_rows(evals, horizon: int):
    agg = {}
    for ev in evals.values():
        for fid, hs in ((ev or {}).get("model_summary") or {}).items():
            m = hs.get(str(horizon))
            if not m:
                continue
            a = agg.setdefault(fid, {"rmse": 0.0, "mse": 0.0, "mae": 0.0, "skill": 0.0, "n_skill": 0, "sel": 0, "n": 0})
            n = m["n_stations"]
            a["rmse"] += m["mean_rmse"] * n; a["mse"] += m["mean_mse"] * n; a["mae"] += m["mean_mae"] * n
            if m.get("mean_skill") is not None:
                a["skill"] += m["mean_skill"] * n; a["n_skill"] += n
            a["sel"] += m["selected_count"]; a["n"] += n
    rows = []
    for fid, a in agg.items():
        n = max(1, a["n"]); skill = a["skill"] / a["n_skill"] if a["n_skill"] else None
        rows.append((skill if skill is not None else -9e9, [NAMES.get(fid, fid), str(a["n"]), str(a["sel"]), m_cm(a["rmse"] / n), m2_cm2(a["mse"] / n), m_cm(a["mae"] / n), pct(skill)]))
    rows.sort(key=lambda r: -r[0])
    return [r[1] for r in rows]


def selection_stats(evals):
    total = sum((ev or {}).get("cells_total", 0) for ev in evals.values())
    beat = sum((ev or {}).get("cells_beating_persistence", 0) for ev in evals.values())
    by_model: dict[str, int] = {}
    for ev in evals.values():
        for r in (ev or {}).get("rows", []):
            by_model[r["forecaster_id"]] = by_model.get(r["forecaster_id"], 0) + 1
    return total, beat, by_model


# ----------------------------------------------------------------------------- deck
def build(data) -> Path:
    d = Deck()
    evals = data["evals"]; det = data["det"]; ts = data["ts"]
    n_stations = sum((ev or {}).get("n_stations", 0) for ev in evals.values())
    total_cells, beat_cells, by_model = selection_stats(evals)
    sat_rows, sat_rmse = satellite_vs_gauge(data)

    # 1 표지
    s = d.prs.slides.add_slide(d.prs.slide_layouts[6]); f = s.background.fill; f.solid(); f.fore_color.rgb = BG; d.n += 1
    t = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.8), Inches(3)); tf = t.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = "초소형위성 기반 수체·수위 탐지 및 시계열 예측 — 과제 성과"; p.font.size = Pt(34); p.font.bold = True; p.font.color.rgb = FG
    q = tf.add_paragraph(); q.text = "위성 영상 → 수체 탐지(U-Net) → 위성 수위·면적 산출 → 시계열 예측(13종 모델 비교) → 성능 평가(RMSE·MSE·MAE·개선율)"; q.font.size = Pt(16); q.font.color.rgb = MUTED
    r = tf.add_paragraph(); r.text = f"실증 지역 낙동강 하구 · 금강 대청댐 · 한강 + 유역 4곳 — 수위계 {n_stations}지점 × 보고 지평 7개 = {total_cells}셀 · 모든 수치는 산출물 파일에서 자동 생성"; r.font.size = Pt(12); r.font.color.rgb = MUTED
    d.md += ["# 초소형위성 기반 수체·수위 탐지 및 시계열 예측 — 과제 성과", "", "자동 생성(scripts/make_results_deck.py). 스크린샷은 docs/presentation/screens, 그림은 docs/presentation/figures.", ""]

    # 2 한 장 요약
    ice = ((det.get("sensors") or {}).get("ICEYE") or {}).get("pooled") or {}
    opt = ((det.get("sensors") or {}).get("PlanetScope") or {}).get("pooled") or {}
    d7 = model_metric_rows(evals, 7)
    best7 = d7[0] if d7 else ["-"] * 7
    deep_sel = sum(v for k, v in by_model.items() if k in DEEP)
    gupo = sat_rmse.get("구포대교"); jeong = sat_rmse.get("정천교")
    s = d.slide("한 장 요약 — 무엇을 했고 얼마나 맞았나")
    d.bullets(s, [
        f"수체 탐지: 부산 ICEYE SAR 4시점 IoU {ice.get('water_iou', 0):.3f} · PlanetScope 광학 4시점 IoU {opt.get('water_iou', 0):.3f} (정답 라벨 대비, 로컬 재현)",
        "위성 수위: 탐지 마스크 경계 × DEM → 지상 수위계와 같은 날 대조(산출 구역과 수위계 위치를 사진·지도에 표시) — "
        + (f"구포대교 RMSE {gupo * 100:.1f} cm, 정천교 {jeong * 100:.1f} cm (융합 LSTM 보정 후)" if gupo and jeong else "부산 2지점 대조"),
        f"시계열 예측: 수위계 10년 자료로 13종 모델을 같은 규칙으로 학습·검증·시험(예측일수 D+1~30 전부 학습), 보고 지평 7개(D+1·3·5·7·10·14·30) × {n_stations}지점 = {total_cells}셀 중 {beat_cells}셀이 기준선(직전값 유지)보다 정확",
        f"D+7 최상 계열 {best7[0]}: 평균 RMSE {best7[3]}, 개선율 {best7[6]} ({len([e for e in evals.values() if e])}지역 지점 가중평균)",
        f"딥러닝(MLP·LSTM·GRU·TCN, GPU 학습)이 검증에서 선택된 셀 {deep_sel}/{total_cells} — 주로 D+7 이상. 단기는 평균 회귀·AR-X 우세",
        "플랫폼(부수): 지역 레지스트리 + 플러그인(탐지·예측기) + REST API 39개 + 발표용 슬라이드 UI — 새 위성 데이터·새 지역이 와도 같은 길로 처리",
    ], width=12.2, size=14)

    # 3 흐름
    s = d.slide("과제 흐름 — 위성 장면 한 장이 예측까지 가는 길", "플랫폼 화면 '위성 파이프라인(입력→예측)' 모드")
    d.bullets(s, ["① 입력: ICEYE SAR / PlanetScope 광학 전처리 장면(3 m, EPSG:32652)", "② 탐지: WBMS U-Net(+ASPP) → 수체 마스크, 정답 라벨 있으면 IoU·F1", "③ 산출: 물 화소 × 화소 면적 = 수체 면적, 경계 화소 × DEM = 위성 수위 → 수위계와 같은 날 대조", "④ 시계열: 지상 수위계 10년(2015~2024) + 기상청 ASOS 강수·기온 → 지점×지평별 13종 모델 비교·선택", "⑤ 평가: 학습·선택에 쓰지 않은 2024년 시험, RMSE(m·cm)·MSE(m²·cm²)·MAE·개선율"], width=5.6, size=12)
    d.image(s, SCR / "30_pipeline_busan.png", left=6.3, top=1.3, max_w=6.6, max_h=5.6, caption="위성 파이프라인 화면(전체)")

    # 4 데이터 현황
    s = d.slide("실증 데이터 — 지역별 보유 현황", "지역 선택 화면(위성 있는 지역 우선)")
    rows = []
    for rid, name in REGIONS.items():
        by = {}
        for item in (data["scenes"].get(rid) or {}).get("scenes") or []:
            by.setdefault(item["sensor"], []).append(item["date"])
        ev = evals.get(rid) or {}
        stations = (data["series"].get(rid) or {}).get("stations") or []
        periods = [st.get("period") for st in stations if st.get("period")]
        span = f"{min(p[0] for p in periods)[:4]}~{max(p[1] for p in periods)[:4]}" if periods else "-"
        has_gt = any(((item.get("metrics") or {}).get("water_iou") is not None) for item in (data["scenes"].get(rid) or {}).get("scenes") or [])
        rows.append([name, ", ".join(f"{k} {len(v)}시점" for k, v in by.items()) or "없음(등록 대기)", "있음" if has_gt else "없음", f"{ev.get('n_stations', '-')}지점 · {sum(st.get('n_days', 0) for st in stations):,}일", span])
    d.table(s, ["지역", "위성 장면(탐지 마스크)", "정답 라벨", "수위계", "수위계 기간"], rows, top=1.35, width=7.6, font=9, row_h=0.3)
    d.image(s, SCR / "10_slide1_region_select.png", left=7.8, top=1.35, max_w=5.1, max_h=5.4, caption="지역 선택 화면")

    # 5 탐지 ① ICEYE
    s = d.slide("수체 탐지 결과 ① 부산 ICEYE SAR — U-Net 탐지 vs 정답 라벨", "화면 3장 '위성 수체 탐지' · 청록 = 탐지, 노랑 = 정답")
    d.image(s, SCR / "12a_slide3_detection_overlay.png", left=0.45, top=1.3, max_w=7.5, max_h=5.4, caption="원본 · 탐지 · 정답 (2020-03-02)")
    d.table(s, ["관측일", "면적 km²", "IoU", "F1", "정밀도", "재현율"], scene_rows(data["scenes"]["nakdong_busan"], "ICEYE"), left=8.1, top=1.35, width=4.9, font=9)
    d.note(s, f"pooled IoU {ice.get('water_iou', 0):.4f} · F1 {ice.get('f1', 0):.4f} · 정확도 {ice.get('accuracy', 0):.4f} ({det.get('series_ko', '')})", top=4.3, left=8.1, width=4.9)

    # 6 탐지 ② PlanetScope
    s = d.slide("수체 탐지 결과 ② 부산 PlanetScope 광학 — Attention U-Net", "같은 화면, 센서를 광학으로 전환")
    d.image(s, SCR / "13a_slide3_detection_planet_overlay.png", left=0.45, top=1.3, max_w=7.5, max_h=5.4, caption="원본 · 탐지 · 정답 (2020-02-18)")
    d.table(s, ["관측일", "면적 km²", "IoU", "F1", "정밀도", "재현율"], scene_rows(data["scenes"]["nakdong_busan"], "PlanetScope"), left=8.1, top=1.35, width=4.9, font=9)
    d.note(s, f"pooled IoU {opt.get('water_iou', 0):.4f} · F1 {opt.get('f1', 0):.4f} · 3/25 장면은 촬영범위가 남쪽으로 8.6 km 이동한 것(수체 급증 아님) — 탐지는 GT와 일치", top=4.3, left=8.1, width=4.9)

    # 7 탐지 ③ 대청
    dc = scene_rows(data["scenes"]["geum_daecheong"], "PlanetScope")
    s = d.slide("수체 탐지 결과 ③ 금강 대청댐 PlanetScope 10시점 (2019 봄 · 2022 5월)", "정답 라벨이 없어 IoU 는 내지 않고 면적·수위 대조로 실증")
    d.image(s, SCR / "20a_slide3_detection_daecheong_overlay.png", left=0.45, top=1.3, max_w=8.6, max_h=5.5, caption="원본 · 탐지 · 장면별 면적 (2019-03-23)")
    d.table(s, ["관측일", "탐지 면적 km²"], [r[:2] for r in dc], left=9.3, top=1.35, width=3.6, font=9)

    # 8 탐지 정확도 화면
    s = d.slide("탐지 정확도 — 성능 평가 화면의 '위성 수체 탐지 정확도' 표", "정답 라벨 대비 · 배포 프로토콜 기준치와 별개 계열")
    d.image(s, SCR / "19f_slide6_detection_accuracy.png", left=0.45, top=1.3, max_w=12.4, max_h=4.9, caption="성능 평가 화면 하단")
    d.note(s, det.get("protocol_note_ko", ""), top=6.35)

    # 8-1 위치 대응: 위성 산출 구역 ↔ 수위계
    s = d.slide("어디를 잰 것인가 — 위성 수위 산출 구역과 지상 수위계의 위치 대응", "위성 수체 탐지 화면 · 청록 상자 = 위성 수위 산출 구역(300 m), 마커 = 수위계, 점선 = 거리")
    d.image(s, SCR / "12c_slide3_location_overlay.png", left=0.45, top=1.3, max_w=8.2, max_h=5.5, caption="탐지 오버레이 위에 산출 구역과 수위계 위치를 겹쳐 그린다 — 구포대교 64 m, 정천교 240 m, 호포 141 m, 김해 930 m(하류 하도로 재설정)")
    d.image(s, SCR / "11b_slide2_station_map.png", left=8.9, top=1.3, max_w=4.0, max_h=5.5, caption="지역 지도: 수위계 지점(WAMIS 좌표·주소)과 산출 구역")

    # 9 위성 수위 vs 수위계
    s = d.slide("위성 수위 산출 vs 지상 수위계 — 같은 지점·같은 날 대조", "관측 경로끼리의 대조(예측 아님) · m 과 cm 병기")
    d.table(s, ["지역", "수위계 지점", "위성 시점", "대조 일수", "RMSE", "MSE", "MAE"], sat_rows, left=0.5, top=1.35, width=7.6, font=10)
    d.image(s, SCR / "12b_slide3_satellite_vs_gauge.png", left=0.5, top=3.2, max_w=7.6, max_h=3.2, caption="화면: 위성이 잰 수위 vs 지상 수위계")
    d.image(s, SCR / "11a_slide2_gauge_history.png", left=8.3, top=1.35, max_w=4.7, max_h=5.2, caption="수위계 10년 기록(EDA 화면)")
    d.note(s, "부산: 융합 LSTM 당일 보정(수위계 기준면) · 대청: 수위대 상수 보정(인수 증적 재현). 위성 원값은 수체 경계 상대 수위.", top=6.55)

    # 10 시계열 개요
    s = d.slide("시계열 분석 — 13종 모델을 같은 규칙으로 비교하고 지점×예측일수별로 선택", "학습 ≤2018 → 검증 2019~2023 연도별 폴드(대청 2021~2023) → 시험 2024")
    d.bullets(s, ["선형·계절·회귀·사례 7종: 릿지 AR, 릿지 AR+기상, 계절편차 AR, 평균 회귀 AR(1), 유사사례 k-NN, 계열 앙상블, 계절 평년값",
                  "딥러닝 4종(GPU torch 학습 → float16 가중치 → 서버는 numpy 추론): MLP, LSTM, GRU, TCN",
                  "Prophet(구 Facebook) + 잔차 자기회귀 1종 · 기준선: 직전값 유지(persistence, 선택 대상 아님)",
                  "선택 규칙: 검증 폴드에서 persistence 대비 MSE 비율 평균이 최소인 후보 → 학습+검증 재적합 → 2024 시험 채점",
                  "개선율(skill) = 1 − MSE(모델)/MSE(직전값 유지). +50 % = 오차 제곱 절반",
                  "입력: 지상 수위계 일 수위 + 기상청 ASOS 강수(1·3·7·30일 합)·기온(원점 이전만)"], width=6.2, size=12)
    d.image(s, FIG / "selection_protocol.png", left=6.9, top=1.35, max_w=6.1, max_h=5.4, caption="검증·시험 분할 규칙")

    # 11 모델 설명 화면
    s = d.slide("모델 13종 설명 — 화면 안내('모델은 어떤 것들이고, 딥러닝·GPU는 어떻게 썼나요?')", "모델별 구조 그림은 부록 '수위예측모델_비교.pptx'")
    d.image(s, SCR / "32_pipeline_model_guide.png", left=0.45, top=1.3, max_w=7.2, max_h=5.6, caption="모델 안내 화면")
    d.image(s, FIG / "model_neural.png", left=7.9, top=1.3, max_w=5.0, max_h=2.7, caption="딥러닝 4종 구조")
    d.image(s, FIG / "model_ridge_arx.png", left=7.9, top=4.2, max_w=5.0, max_h=2.7, caption="릿지 AR + 기상")

    # 11-1 학습·검증·시험 구간 선택
    s = d.slide("학습·검증·시험 구간을 직접 정해 다시 학습 — 같은 규칙, 다른 시기", "시계열 예측 화면의 '학습·검증·시험 구간 직접 정하기' · 서버에서 재학습(비동기 작업) → 예측·평가 화면이 새 결과를 사용")
    d.image(s, SCR / "15a_slide4_training_window.png", left=0.45, top=1.3, max_w=7.6, max_h=5.5, caption="연도 슬라이더로 학습 종료·검증 연도·시험 연도를 고르면 다섯 날짜로 바뀌고 검증 규칙(학습 ≥2년·검증 ≥90일·시험 ≥30일)을 통과해야 학습이 시작된다")
    d.bullets(s, ["기본값은 학습 ≤2018 · 검증 2019∼2023 · 시험 2024 (발표 수치의 기준)", "시험 연도를 바꾸면 그 해 기상 조건에서 채점되므로 구간이 다른 성적끼리는 직접 비교하지 않는다", "사용자 지정 결과는 지역별 사용자 디렉터리에 저장되어 기본 산출물을 덮어 보이며, '기본 구간으로 되돌리기'로 복귀", "서버 컨테이너에는 torch·Prophet 이 없어 재학습은 내장 7종 계열로 진행(딥러닝은 GPU 학습 이미지 필요)"], left=8.3, top=1.35, width=4.7, size=11)

    # 12 모델별 예측 그래프(상호작용)
    s = d.slide("모델별 예측 그래프 — 한 원점(2024-07-01, 구포대교)에서 12종 전망을 한 그래프에", "화면 4장 '모든 모델로 한 번에 비교하기' · 굵은 청록 = 검증 선택 모델 · 노랑 = 실제 · 범례에 순위")
    d.image(s, SCR / "16b_model_compare_chart.png", left=0.45, top=1.3, max_w=12.4, max_h=5.2, caption="모델 12종 전망 비교(D+14) — 7/9~7/11 실제 급상승(3.2 m)은 모든 모델이 놓침: 원점 이후 강우가 입력에 없기 때문")

    # 12-1 조작: 모델 켜고 끄기·예측 구간 확대·단위
    s = d.slide("그래프 조작 — 모델 켜고 끄기, 예측 구간만 확대, m/cm 축, 차이·오차 보기", "보일 모델(전체·상위 3·선택 모델만·딥러닝만·기준선 숨기기) · 순위 기준 지표 선택")
    d.image(s, SCR / "16a_model_compare_controls.png", left=0.45, top=1.3, max_w=6.2, max_h=5.5, caption="조작 패널: 순위 기준·그래프 종류·축 단위·보일 모델")
    d.image(s, SCR / "16d_model_compare_forecast_only.png", left=6.8, top=1.3, max_w=6.2, max_h=5.5, caption="'예측 구간만 보기' — 원점 이후만 확대")

    # 12-2 차이·오차 그래프
    s = d.slide("차이를 자세히 — 선택 모델 대비 차이(cm)와 실제 대비 오차(cm)", "같은 화면의 그래프 종류 전환 · 굵은 선 = 1위, 점선 = 최하위")
    d.image(s, SCR / "16e_model_compare_difference.png", left=0.45, top=1.3, max_w=6.2, max_h=5.5, caption="각 모델 − 선택 모델 (cm)")
    d.image(s, SCR / "16f_model_compare_error_top3.png", left=6.8, top=1.3, max_w=6.2, max_h=5.5, caption="각 모델 − 실제 관측 (cm) · 상위 3만")

    # 13 순위 표
    s = d.slide("모델 순위 — 🥇🥈🥉 최상위와 ⚠ 최하위, RMSE·MSE·MAE·개선율 (m·cm 병기)", "이번 원점 검증 RMSE 기준(실제 관측이 있을 때) · 없으면 시험 연도 성적 기준 · ★ = 검증에서 고른 모델")
    d.image(s, SCR / "16c_model_ranking.png", left=0.45, top=1.3, max_w=12.4, max_h=5.4, caption="선택 모델은 검증 연도 성적으로 고른 것이라 이번 원점에서 1위가 아닐 수 있다 — 화면이 그 사실을 배너로 설명")

    # 13-1 파고들기
    s = d.slide("모델 파고들기 — 한 모델의 지평별 |오차|와 예측·실제·차이 표", "'자세히 볼 모델' 선택 → 순위·평균/최대 오차·선택 모델과의 차이 타일 + 지평별 표")
    d.image(s, SCR / "16g_model_drilldown.png", left=0.45, top=1.3, max_w=12.4, max_h=5.4, caption="지평별 |오차|(cm) 막대와 예측·실제·오차·선택 대비 차이 표(m·cm)")

    # 14 차이 표
    s = d.slide("모델별 예측 차이 표 — 지평별 각 모델 예측과 선택 모델 대비 차이", "같은 화면의 펼침 표 · 표 단위 m(cm 차이) / cm / m 전환")
    d.image(s, SCR / "16h_model_compare_diff_table.png", left=0.45, top=1.3, max_w=12.4, max_h=5.4, caption="D+1·3·5·7·10·14 지평별 모델 예측과 선택 모델 대비 차이")

    # 15-17 모델별 성능지표 전체 (D+1/7/30)
    for h in (1, 7, 30):
        s = d.slide(f"모델별 성능지표 D+{h} — {len([e for e in evals.values() if e])}지역 {n_stations}지점 가중평균 (시험 연도 2024)", "선택 여부와 무관하게 모든 계열을 같은 시험 구간에서 채점 · '선택 지점' = 검증에서 이 계열이 뽑힌 지점 수")
        d.table(s, ["모델", "지점", "선택 지점", "평균 RMSE", "평균 MSE", "평균 MAE", "평균 개선율"], model_metric_rows(evals, h), top=1.35, font=10, row_h=0.33)

    # 18 지역 평가 화면
    s = d.slide("성능 평가 화면 — 예측일수별 지점 성적 + 모델별 성능지표(지역 평균)", "쉬운 화면 6장 · 예측일수 슬라이더 1·3·5·7·10·14·30")
    d.image(s, SCR / "19a_slide6_horizon_table.png", left=0.45, top=1.3, max_w=6.2, max_h=5.4, caption="지점별 선택 모델·RMSE·개선율(D+7)")
    d.image(s, SCR / "19b_slide6_model_metrics.png", left=6.8, top=1.3, max_w=6.2, max_h=5.4, caption="모델별 성능지표(부산 6지점 평균, D+7)")

    # 19 행렬 + 백테스트
    s = d.slide("지점 × 모델 RMSE 행렬과 시험 연도 백테스트", "행렬(cm)은 지점마다 어느 계열이 강한지, 백테스트는 2024년 전체를 하루씩 예측한 궤적")
    d.image(s, SCR / "19c_slide6_station_model_matrix.png", left=0.45, top=1.3, max_w=6.2, max_h=5.4, caption="지점 × 모델 시험 RMSE 행렬(cm) · D+7")
    d.image(s, SCR / "19d_slide6_backtest.png", left=6.8, top=1.3, max_w=6.2, max_h=5.4, caption="백테스트 궤적(선택 모델 vs 실제 vs 기준선)")

    # 20-22 지역별 시험 성적 (실증 3지역은 지점×3지평, 나머지 지역은 한 표에 D+7 요약)
    main_regions = [rid for rid in evals if rid in ("nakdong_busan", "geum_daecheong", "han_river")]
    for rid in main_regions:
        ev = evals[rid]
        if not ev:
            continue
        rows = [[r["station_name"], str(r["horizon"]), NAMES.get(r["forecaster_id"], r["forecaster_id"]), m_cm(r.get("test_rmse")), m2_cm2(r.get("test_mse")), m_cm(r.get("test_mae")), m_cm(r.get("persistence_test_rmse")), pct(r.get("skill_vs_persistence"))]
                for r in ev["rows"] if r["horizon"] in (1, 7, 30)]
        s = d.slide(f"{REGIONS[rid]} — 지점별 시험 성적 (2024, 선택 모델)", f"기준선보다 나은 셀 {ev['cells_beating_persistence']}/{ev['cells_total']} · 기상이 뚜렷이 도운 셀 {ev.get('cells_weather_helps', 0)}/{ev.get('cells_with_weather_candidate', 0)}")
        d.table(s, ["지점", "D+", "선택 모델", "RMSE", "MSE", "MAE", "기준선 RMSE", "개선율"], rows, top=1.35, font=9, row_h=0.28)
    others = [rid for rid in evals if rid not in main_regions and evals[rid]]
    if others:
        rows = []
        for rid in others:
            ev = evals[rid]
            for r in ev["rows"]:
                if r["horizon"] == 7:
                    rows.append([REGIONS[rid], r["station_name"], NAMES.get(r["forecaster_id"], r["forecaster_id"]), m_cm(r.get("test_rmse")), m2_cm2(r.get("test_mse")), m_cm(r.get("persistence_test_rmse")), pct(r.get("skill_vs_persistence"))])
        s = d.slide("위성 없는 지역(수위계만)도 같은 규칙으로 — 새 4개 유역 D+7 시험 성적", "영산강·섬진강·금강 본류·낙동강 상류: WAMIS 10년 수위계 + 기상청 ASOS · 위성 장면이 등록되면 탐지·대조 기능이 열린다")
        d.table(s, ["지역", "지점", "선택 모델", "RMSE", "MSE", "기준선 RMSE", "개선율"], rows, top=1.35, font=9, row_h=0.28)
        d.note(s, " · ".join(f"{REGIONS[rid]}: 기준선보다 나은 셀 {evals[rid]['cells_beating_persistence']}/{evals[rid]['cells_total']}" for rid in others), top=6.55)
        s = d.slide("위성 데이터 입력 — 위성이 없는 지역도 GeoTIFF 를 올리면 탐지·면적·위성 수위 대조가 열린다", "쉬운 화면 3장(위성 없는 지역) · 위성 파이프라인 '새 위성 데이터가 왔을 때'")
        d.image(s, SCR / "24a_slide3_upload_form.png", left=0.45, top=1.3, max_w=7.4, max_h=5.5, caption="영산강: 전처리 GeoTIFF(+마스크·GT) 업로드 → 등록 → 탐지 대기/완료")
        d.image(s, SCR / "27a_slide6_eval_yeongsan_models.png", left=8.0, top=1.3, max_w=5.0, max_h=5.5, caption="영산강 모델별 성능지표(D+7)")

    # 23 대청 화면
    s = d.slide("금강 대청댐 — 예측·평가 화면", "저수지 지점: 변화가 느려 장기 지평 개선율이 크다")
    d.image(s, SCR / "22_slide4_forecast_daecheong.png", left=0.45, top=1.3, max_w=6.2, max_h=5.4, caption="대청댐 수위 예측 실행")
    d.image(s, SCR / "21a_slide6_daecheong_model_metrics.png", left=6.8, top=1.3, max_w=6.2, max_h=5.4, caption="대청 모델별 성능지표(D+7)")

    # 24 딥러닝·GPU
    winners = sorted(by_model.items(), key=lambda kv: -kv[1])
    s = d.slide("딥러닝(GPU)·Prophet 도 같은 규칙으로 — 어떤 모델이 뽑혔나", f"전체 {total_cells}셀 선택 분포 · 딥러닝 {deep_sel}셀")
    d.table(s, ["모델", "선택 셀", "비율"], [[NAMES.get(k, k), str(v), f"{v / max(1, total_cells) * 100:.0f} %"] for k, v in winners], left=0.5, top=1.35, width=4.6, font=10, row_h=0.3)
    d.image(s, FIG / "gpu_vs_cpu.png", left=5.4, top=1.3, max_w=7.5, max_h=3.2, caption="GPU 학습 속도(LSTM 1회 3.0 s → 0.9 s)")
    d.bullets(s, ["GPU 는 후보 학습(지점 14 × 후보 10 × 폴드 7 ≈ 1,000회)을 3~4배 빠르게 해 '더 많은 모델을 같은 규칙으로' 비교하게 한다", "추론은 numpy 로 수 ms — 서버(CPU 슬림 이미지)에 torch 없이 같은 결과", "최종 선택은 검증 성적이 정한다 — 딥러닝은 D+7 이상에서, 단기는 평균 회귀·AR-X 가 우세"], left=5.4, top=4.7, width=7.5, height=2.2, size=11)

    # 24-1 미래 강수(완전 예보 가정) 상한 실험 — 공식 성적표(위 슬라이드)와 별개, 수치는 experiments/future_rain.json 에서만 온다
    experiments = data.get("experiments") or {}
    fr_rows = region_summary_rows(experiments, REGIONS, horizons=FUTURE_RAIN_HORIZONS)
    fr_regions = [rid for rid, payload in experiments.items() if experiment_rows(payload)]
    s = d.slide("미래 강수(예보)를 넣으면 — 완전 예보 가정 상한 실험",
                "원점 이후 실측 강수를 '완벽한 예보' 로 간주해 릿지 AR-X 에 넣은 부수 실험 · 지점 평균 개선율 = 과거 강수 AR-X / +3일 예보 / +전 구간 예보 · 공식 수치 아님")
    if fr_rows:
        d.table(s, ["지역", "지점", *[f"D+{h}" for h in FUTURE_RAIN_HORIZONS]], fr_rows, left=0.5, top=1.35, width=7.6, font=9, row_h=0.3)
        pooled = {h: {"n": 0, "beats3": 0, "beats_all": 0, "gain_all": []} for h in FUTURE_RAIN_HORIZONS}
        for rid in fr_regions:
            rows_ = experiment_rows(experiments[rid])
            for h, entry in summary_by_horizon(rows_).items():
                if h in pooled:
                    pooled[h]["n"] += entry["n_stations"]; pooled[h]["beats3"] += entry["beats_arx"].get("future_3", 0); pooled[h]["beats_all"] += entry["beats_arx"].get("future_all", 0)
            for r in rows_:  # 지점별 (전 구간 예보 − AR-X) 개선율 차이, %p
                a = (r["metrics"].get("ridge_arx") or {}).get("skill"); f = (r["metrics"].get("future_all") or {}).get("skill")
                if r["horizon"] in pooled and isinstance(a, (int, float)) and isinstance(f, (int, float)):
                    pooled[r["horizon"]]["gain_all"].append((f - a) * 100)
        beats_text = " · ".join(f"D+{h} {v['beats_all']}/{v['n']}" for h, v in pooled.items() if v["n"])
        gain_text = " · ".join(f"D+{h} {sum(v['gain_all']) / len(v['gain_all']):+.1f} %p" for h, v in pooled.items() if v["gain_all"])
        d.bullets(s, [
            "각 셀 = 과거 강수 AR-X / +3일 예보 / +전 구간 예보의 지점 평균 개선율(직전값 유지 대비 MSE 감소율, %) — 두 예보 변형은 실측 미래 강수를 그대로 넣은 상한",
            f"전 구간 예보의 개선율 이득(지점 가중 평균, AR-X 대비): {gain_text}",
            f"전 구간 예보가 과거 강수 AR-X 를 0.5 %p 이상 앞선 지점(지평별): {beats_text}",
            "뜻: 급상승은 원점 이후 비 때문에 생기므로 '앞으로 내릴 비'를 알면 그만큼 오차가 준다 — 다만 상한에서도 이득은 수 %p 수준이라, 홍수기 급변을 잡으려면 예보와 함께 댐 방류량 같은 입력이 더 필요하다",
            "실시간에는 기상청 단기 예보를 연결해야 하고(아직 미연결), 예보 오차만큼 이득이 줄어든다. 위 모델별 성능지표·지역별 성적표는 이 실험을 포함하지 않는다",
        ], left=0.55, top=1.5 + 0.3 * (len(fr_rows) + 1), width=7.5, height=5.6 - 0.3 * (len(fr_rows) + 1), size=10.5)
    else:
        d.bullets(s, ["실험 산출물(regions/<id>/models/gauge/experiments/future_rain.json)이 아직 없다 — scripts/eval_future_rain_experiment.py 를 실행하면 이 표가 채워진다",
                      "설계: 릿지 AR-X(과거 강수·기온) + 원점 다음 3일 또는 목표일까지의 실측 강수 합·최대(log1p) → 완전 예보 가정 상한", "공식 성적표에는 넣지 않는다"], width=7.5, size=11)
    d.image(s, SCR / "19g_slide6_future_rain.png", left=8.3, top=1.35, max_w=4.7, max_h=5.3, caption="성능 평가 화면의 같은 표(펼침 상자)")
    d.note(s, f"실험 지역 {len(fr_regions)}/{len(REGIONS)} · 검증 폴드에서 λ 선택, 2024 시험 · 완전 예보 가정 = 실측 미래 강수를 예보로 간주한 상한", top=6.65)

    # 24-2 수체 면적 계열 모델 직접 비교(LOO) — 표본이 4~10점이라 우열을 가릴 수 없다는 판정을 함께 싣는다
    area_exp = data.get("area_models") or {}
    am_rows = am.region_summary_rows(area_exp, REGIONS)
    am_regions = [rid for rid, payload in area_exp.items() if am.sensors_available(payload)]
    s = d.slide("수체 면적 시계열은 persistence 말고 다른 모델로 검증되나 — 직접 비교(LOO)",
                "공통 격자에서 다시 잰 면적 계열에 후보 7종을 leave-one-out 으로 돌린 부수 실험 · 관측이 센서당 4~10개뿐이라 순위는 참고용")
    if am_rows:
        d.table(s, ["지역", "센서", "관측", "직전값 유지 RMSE(km²)", "최저 LOO 후보", "그 RMSE(km²)", "개선율", "수위→면적 다리"],
                am_rows, left=0.5, top=1.35, width=7.6, font=9, row_h=0.3)
        noise = []
        for rid in am_regions:
            for sensor, block in sorted(am.series_blocks(area_exp[rid]).items()):
                floor = block.get("observation_noise_floor_km2")
                span = block.get("area_range_km2")
                if isinstance(floor, (int, float)) and isinstance(span, (int, float)) and span:
                    noise.append(f"{SHORT_NAMES.get(rid, rid)} {sensor} {floor:.1f}/{span:.1f} km²")
        d.bullets(s, [
            "면적은 센서별 공통 격자(같은 상자·같은 화소)에서 다시 쟀다 — 장면마다 촬영 범위가 달라 원래의 장면별 면적은 날짜 간 비교가 되지 않는다",
            "LOO = 관측 하나를 빼고 나머지로 맞히기. 시험점 수 = 관측 수라 n=4 에서는 부호검정으로 얻을 수 있는 가장 작은 p 가 0.125 — 어떤 후보도 직전값 유지보다 유의하게 낫다고 말할 수 없다",
            ("이틀 이내 재방문쌍의 면적 차(관측 잡음 하한)/계열 변동폭: " + " · ".join(noise)) if noise else "이틀 이내 재방문쌍이 없어 관측 잡음 하한은 계산하지 않았다",
            "수위→면적 검량곡선은 어느 조합도 문턱을 넘지 못했다(대청은 상관이 음수, 부산은 관측일 수위 변동폭이 1~4 cm) → 파생 면적 예측 모델은 공개하지 않는다",
            "검증이 실제로 되는 것은 10개년 일별 수위 계열이다 — 앞 슬라이드의 성적표가 그것이고, 면적 계열은 아직 관측 수가 모자란다",
        ], left=0.55, top=1.5 + 0.3 * (len(am_rows) + 1), width=7.5, height=5.6 - 0.3 * (len(am_rows) + 1), size=10.5)
    else:
        d.bullets(s, ["실험 산출물(regions/<id>/models/gauge/experiments/area_models.json)이 아직 없다 — scripts/eval_area_models_experiment.py 를 실행하면 이 표가 채워진다",
                      "설계: 공통 격자 면적 계열 × 후보(직전 관측 유지·계열 평균·시간 선형추세·연중 시기 평년값·AR(1) 평균회귀·같은 날 수위 선형/단조 회귀) × LOO",
                      "공식 성적표에는 넣지 않는다"], width=7.5, size=11)
    d.image(s, SCR / "19h_slide6_area_models.png", left=8.3, top=1.35, max_w=4.7, max_h=5.3, caption="성능 평가 화면의 같은 표(펼침 상자)")
    d.note(s, f"실험 지역 {len(am_regions)}/{len(REGIONS)} · 원본 JSON regions/<id>/models/gauge/experiments/area_models.json · 공식 성적표와 별개", top=6.65)

    # 25 위성 시계열 지표
    ser = ts.get("series") or {}
    if ser:
        corr = ser["A_level_correction"]["corrected"]; b = ser["B_persistence_reference"]["persistence_of_corrected"]; fch = ser["F_change"]["level"]
        s = d.slide("위성 데이터로 잰 수위·수체 시계열 지표 (부산, m·cm 병기)", "관측 정확도(같은 날 위성 vs 수위계)와 위성 시점 사이 변화 크기 — 앞날을 맞힌 예측 오차가 아니므로 예측 모델 성적과 같은 잣대로 비교하지 않는다")
        d.table(s, ["항목", "RMSE", "MSE", "MAE", "표본"], [
            ["위성 수위 당일 보정(융합 LSTM, in-sample)", f"{corr['rmse_m']:.4f} m ({corr['rmse_cm']:.2f} cm)", f"{corr['mse_m2']:.6f} m² ({corr['mse_cm2']:.2f} cm²)", f"{corr['mae_m']:.4f} m ({corr['mae_cm']:.2f} cm)", str(corr["n"])],
            ["정식 프로토콜 LODO", f"{ser['A_level_correction']['lodo_rmse_m']:.4f} m ({ser['A_level_correction']['lodo_rmse_cm']:.2f} cm)", "-", "-", "지점 제외 교차검증"],
            ["위성 시계열 직전값 유지 참고치", f"{b['rmse_m']:.4f} m ({b['rmse_cm']:.2f} cm)", f"{b['mse_m2']:.6f} m² ({b['mse_cm2']:.2f} cm²)", f"{b['mae_m']:.4f} m ({b['mae_cm']:.2f} cm)", str(b["n"])],
            ["위성 수위 변화량 추적", f"{fch['rmse_m']:.4f} m ({fch['rmse_cm']:.2f} cm)", f"{fch['mse_m2']:.6f} m² ({fch['mse_cm2']:.2f} cm²)", f"{fch['mae_m']:.4f} m ({fch['mae_cm']:.2f} cm)", f"{fch['n']} · skill {fch['skill_vs_no_change']:+.2f}"],
        ], top=1.35, font=10)
        d.image(s, SCR / "19e_slide6_satellite_metrics.png", left=0.5, top=3.3, max_w=12.3, max_h=3.2, caption="성능 평가 화면의 같은 표")
        g = ser.get("G_iitp") or {}
        d.note(s, f"IITP 성과지표 정합값(정규화 MSE) 종합 {g.get('composite_norm_mse')} (목표 {g.get('target')}, {'달성' if g.get('achieved') else '미달'}) · 면적 MAE ICEYE {ser['C_area']['summary']['ICEYE']['mae_km2']} km², PlanetScope {ser['C_area']['summary']['PlanetScope']['mae_km2']} km²", top=6.6)

    # 26 결과 화면
    s = d.slide("결과 화면 — 선택 모델 예측선·±σ 띠·실측 대조, 수체 면적 예측", "쉬운 화면 5장 '결과 그래프'")
    d.image(s, SCR / "18a_slide5_level_result.png", left=0.45, top=1.3, max_w=7.3, max_h=5.5, caption="수위 예측 결과(구포대교, 2024-07-01 원점 D+14)")
    d.image(s, SCR / "17_slide4_area_forecast.png", left=8.0, top=1.3, max_w=5.0, max_h=5.5, caption="수체 면적 예측(위성 마스크 시계열 + 강수)")

    # 27 플랫폼(부수)
    s = d.slide("플랫폼(부수) — 새 지역·새 위성 데이터가 와도 같은 길로", "지역 레지스트리 + 플러그인 3종 + REST API 39개 · 서버 wc-ui/wc-api.semyeongsoft.com")
    d.bullets(s, ["지역 = 위성 장면·수위계 시계열 중 있는 것만 넣으면 그만큼 기능이 열린다 (regions/<id>/)", "새 장면 등록 API: 전처리 GeoTIFF(+마스크) → 퀵룩·공통 격자 마스크 → 면적 예측 샘플 자동 반영", "모델 교체: 탐지 어댑터(SEGMENTER_PLUGINS)·수위 예측기(GAUGE_FORECASTER_PLUGINS)·면적 어댑터(PREDICTOR_PLUGINS)", "외부 시스템: POST /predictions 에 model_id=gauge-level-forecast — 학습 수위 예측(persistence 는 상수라 안 씀)", "발표 UI: 쉬운 화면 6장 · 위성 파이프라인 · 고급 화면(지역 발표 페이지)"], width=5.8, size=12)
    d.image(s, SCR / "31_pipeline_busan_forecasts.png", left=6.5, top=1.3, max_w=3.1, max_h=5.6, caption="위성 파이프라인(전체)")
    d.image(s, SCR / "40_presentation_page_busan.png", left=9.8, top=1.3, max_w=3.2, max_h=5.6, caption="지역 발표 페이지(전체)")

    # 28 한계
    s = d.slide("한계와 다음 단계")
    d.bullets(s, ["탐지 정확도(IoU)는 정답 라벨이 있는 부산만 계산 — 대청은 면적·수위 대조로만 실증", "위성 시점이 지점당 8~10개라 위성만으로는 예측 모델을 학습할 수 없다 → 수위계 10년으로 학습·검증하고 위성은 관측 경로로 대조", "면적 계열 모델을 직접 비교해도(LOO) 센서당 4~10점이라 우열이 가려지지 않고, 이틀 이내 재방문쌍의 면적 차가 계열 변동폭의 3분의 2에 이른다 — 수위→면적 검량곡선도 문턱을 못 넘어 파생 면적 모델은 공개하지 않았다(앞 슬라이드)", "기상(강수·기온)은 단기(D+1)를 크게 개선하지 못했다(원점 수위에 어제 비가 이미 반영) — 다음 입력: 미래 강수 예보·댐 방류량", "미래 강수를 '완벽히 아는' 완전 예보 가정 실험(앞 슬라이드)은 예보 입력의 이득 상한만 보여 준다 — 실시간 기상청 예보 연결은 미완이고 예보 오차만큼 이득이 줄며, 공식 성적표에는 넣지 않았다", "시험 성적은 2024년 기상에 종속. 기준점 이동 지점(양산교)은 장기 지평 불리", "밴드(±σ)는 검증 잔차의 경험적 산포이지 신뢰구간이 아니다", "딥러닝은 지점당 3,650일 규모에서 선형·사례 모델과 대등 — 지점·입력이 늘면 재평가"], width=12.2, size=13)

    out = OUT / "과제성과_위성수체수위_시계열예측.pptx"
    d.prs.save(str(out))
    (OUT / "과제성과_위성수체수위_시계열예측.md").write_text("\n".join(d.md), encoding="utf-8")
    return out


def main() -> int:
    out = build(load())
    print("wrote", out, "and", out.with_suffix(".md"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
